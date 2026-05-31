from pptx.util import Inches, Pt

from deckbridge.renderers.common.style_resolver import resolve_text_style
from deckbridge.renderers.gslides.utils import GSLIDES_ALIGN_MAP, GSLIDES_VERTICAL_ALIGN_MAP, hex_to_slides_rgb, inches_to_emu
from deckbridge.renderers.pptx.utils import PPTX_ALIGN_MAP, PPTX_VERTICAL_ALIGN_MAP, hex_to_rgb255


def resolve_text_content(slide, slot_key, slot):
    """Resolve the text content for a given slot.

    This helper extracts the appropriate text based on the slot's ``content_type``.
    For ``chart_title`` slots it composes a list of lines containing the chart
    title and optional subtitle, along with any style overrides. For other
    content types it returns the raw text stored under ``slot_key`` in the slide.

    Args:
        slide (dict): The current slide definition containing a ``content``
            mapping of block identifiers to block objects.
        slot_key (str): The key that identifies the slot within the slide layout.
        slot (dict): The slot definition, which may include ``content_type`` and
            ``source`` fields.

    Returns:
        tuple: ``(text, style_overrides)`` where ``text`` is either a string or a
        list of ``{"text": str, "style_key": str}`` dictionaries, and
        ``style_overrides`` is the block's ``style_overrides`` (or ``None``).
    """
    content_type = slot.get("content_type")

    if content_type == "chart_title":
        block = slide["content"].get(slot["source"])
        if not block:
            return None, None

        style_overrides = block.style_overrides

        title = block.chart_title
        subtitle = getattr(block, "chart_subtitle", None)

        lines = []

        if title:
            lines.append(
                {
                    "text": title + ("\n" if subtitle else ""),
                    "style_key": "chart_title",
                }
            )

        if subtitle:
            lines.append(
                {
                    "text": subtitle,
                    "style_key": "chart_subtitle",
                }
            )

        return (lines, style_overrides) if lines else (None, None)

    return slide.get(slot_key), None


def render_text_slot(ctx, slot, text, slot_key, style_overrides):
    """Render a text slot using the appropriate backend.

    Dispatches to the backend‑specific rendering functions for PPTX or Google Slides.
    If ``text`` is falsy, the function returns without performing any rendering.

    Args:
        ctx: Rendering context containing backend information and theme.
        slot: Slot definition dictionary from the layout.
        text: Text content to render; can be a string or a list of dictionaries with
            ``"text"`` and ``"style_key"`` keys.
        slot_key: Identifier for the slot used for style resolution.
        style_overrides: Optional style overrides dictionary.

    Returns:
        None. The rendering functions modify the context or slide objects in‑place.
    """
    if not text:
        return

    if ctx.backend == "pptx":
        _render_text_pptx(ctx, slot, text, slot_key, style_overrides)

    elif ctx.backend == "gslides":
        _render_text_gslides(ctx, slot, text, slot_key, style_overrides)

    else:
        raise ValueError(f"Unsupported backend: {ctx.backend}")


def _render_text_pptx(ctx, slot, text, slot_key, style_overrides):
    """Render a text box onto a PPTX slide.

    Creates a textbox shape on the current PPTX slide and populates it with the
    supplied ``text``. The function resolves style information for each line of
    text and applies paragraph‑level alignment.

    Args:
        ctx: Rendering context containing the current PPTX slide object and
            theme information.
        slot: Dictionary describing the slot dimensions (x, y, w, h) and other
            layout metadata.
        text: Text content to render; can be a string or a list of dictionaries
            with ``"text"`` and ``"style_key"`` keys.
        slot_key: Identifier used for style resolution of the base slot.
        style_overrides: Optional dictionary of style overrides that supersede
            theme defaults.

    Returns:
        None. The PPTX slide is modified in‑place.
    """
    textbox = ctx.slide_obj.shapes.add_textbox(
        Inches(slot["x"]),
        Inches(slot["y"]),
        Inches(slot["w"]),
        Inches(slot["h"]),
    )

    tf = textbox.text_frame
    tf.clear()

    base_style = resolve_text_style(slot_key, slot, ctx.theme, ctx.layout_spec.name, style_overrides)
    tf.vertical_anchor = PPTX_VERTICAL_ALIGN_MAP[base_style.get("vertical_align", "TOP")]

    p = tf.paragraphs[0]

    lines = text if isinstance(text, list) else [{"text": text, "style_key": slot_key}]

    for line in lines:
        run = p.add_run()
        run.text = line["text"]

        style = resolve_text_style(line["style_key"], {"style_key": line["style_key"]}, ctx.theme, ctx.layout_spec.name, style_overrides)

        run.font.size = Pt(style["font_size"])
        run.font.bold = style["bold"]
        run.font.italic = style["italic"]
        run.font.underline = style["underline"]
        run.font.color.rgb = hex_to_rgb255(style["font_color"])

    # alignment applies at paragraph level
    p.alignment = PPTX_ALIGN_MAP[base_style["align"]]


def _render_text_gslides(ctx, slot, text, slot_key, style_overrides):
    """Render a text box onto a Google Slides page.

    Constructs a ``TEXT_BOX`` shape in the Google Slides API request list and
    populates it with the supplied ``text``. Styles are applied per text range
    based on resolved style information, and paragraph alignment is set for the
    entire box.

    Args:
        ctx: Rendering context that supplies the current page ID, theme, and
            layout specifications.
        slot: Dictionary describing the slot dimensions (x, y, w, h) and other
            layout metadata.
        text: Text content to render; can be a string or a list of dictionaries
            with ``"text"`` and ``"style_key"`` keys.
        slot_key: Identifier used for style resolution of the base slot.
        style_overrides: Optional dictionary of style overrides that supersede
            theme defaults.

    Returns:
        None. The function adds the necessary API requests to ``ctx`` which are
        later executed when the slide is rendered.
    """
    object_id = f"{slot_key}_{ctx.page_id}"

    requests = []

    lines = text if isinstance(text, list) else [{"text": text, "style_key": slot_key}]

    # -----------------------
    # Build full text + ranges
    # -----------------------
    full_text = ""
    ranges = []

    cursor = 0

    for line in lines:
        line_text = line["text"]
        start = cursor
        end = cursor + len(line_text)

        ranges.append((start, end, line["style_key"]))
        full_text += line_text
        cursor = end

    # -----------------------
    # Create box
    # -----------------------
    requests.append(
        {
            "createShape": {
                "objectId": object_id,
                "shapeType": "TEXT_BOX",
                "elementProperties": {
                    "pageObjectId": ctx.page_id,
                    "size": {
                        "height": {"magnitude": inches_to_emu(slot["h"]), "unit": "EMU"},
                        "width": {"magnitude": inches_to_emu(slot["w"]), "unit": "EMU"},
                    },
                    "transform": {
                        "scaleX": 1,
                        "scaleY": 1,
                        "translateX": inches_to_emu(slot["x"]),
                        "translateY": inches_to_emu(slot["y"]),
                        "unit": "EMU",
                    },
                },
            }
        }
    )

    # -----------------------
    # Insert full text
    # -----------------------
    requests.append(
        {
            "insertText": {
                "objectId": object_id,
                "text": full_text,
            }
        }
    )

    # -----------------------
    # Apply styles per range
    # -----------------------
    for start, end, style_key in ranges:
        style = resolve_text_style(style_key, {"style_key": style_key}, ctx.theme, ctx.layout_spec.name, style_overrides)

        api_style = {
            "fontSize": {"magnitude": style["font_size"], "unit": "PT"},
            "foregroundColor": {"opaqueColor": {"rgbColor": hex_to_slides_rgb(style["font_color"])}},
            "bold": style["bold"],
            "italic": style["italic"],
            "underline": style["underline"],
        }

        requests.append(
            {
                "updateTextStyle": {
                    "objectId": object_id,
                    "textRange": {
                        "type": "FIXED_RANGE",
                        "startIndex": start,
                        "endIndex": end,
                    },
                    "style": api_style,
                    "fields": ",".join(api_style.keys()),
                }
            }
        )

    # -----------------------
    # Alignment (whole paragraph)
    # -----------------------
    base_style = resolve_text_style(slot_key, slot, ctx.theme, ctx.layout_spec.name, style_overrides)

    requests.append(
        {
            "updateParagraphStyle": {
                "objectId": object_id,
                "textRange": {"type": "ALL"},
                "style": {"alignment": GSLIDES_ALIGN_MAP[base_style["align"]]},
                "fields": "alignment",
            }
        }
    )

    vertical_align = base_style.get("vertical_align", "TOP")

    requests.append(
        {
            "updateShapeProperties": {
                "objectId": object_id,
                "shapeProperties": {
                    "contentAlignment": GSLIDES_VERTICAL_ALIGN_MAP[vertical_align],
                },
                "fields": "contentAlignment",
            }
        }
    )

    ctx.add_create_text_requests(requests)
